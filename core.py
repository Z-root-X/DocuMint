import os
import re
import time
import pandas as pd
from docx import Document
import win32com.client
from datetime import datetime

def replace_placeholders_in_doc(doc, replacements):
    """Replaces placeholders in a .docx document.

    This function iterates through all paragraphs and tables in a document
    and replaces placeholder strings (e.g., `<Name>`) with their corresponding
    values from the replacements dictionary.

    Args:
        doc (docx.Document): The python-docx Document object.
        replacements (dict): A dictionary where keys are placeholders
            and values are the replacement strings.
    """
    for para in doc.paragraphs:
        full_text = "".join(run.text for run in para.runs)
        new_text = full_text
        for key, val in replacements.items():
            new_text = new_text.replace(key, str(val))
        if new_text != full_text:
            p = para._p
            for child in list(p):
                p.remove(child)
            para.add_run(new_text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    full_text = "".join(run.text for run in para.runs)
                    new_text = full_text
                    for key, val in replacements.items():
                        new_text = new_text.replace(key, str(val))
                    if new_text != full_text:
                        p = para._p
                        for child in list(p):
                            p.remove(child)
                        para.add_run(new_text)

def is_valid_email(email_address):
    """Validates an email address using a regular expression.

    Args:
        email_address (str): The email address to validate.

    Returns:
        bool: True if the email address is valid, False otherwise.
    """
    pattern = r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"
    return re.match(pattern, email_address) is not None

def process_emails(data_file, template_file, pdf_folder, logs_folder, log_callback, 
                   email_subject, email_body, pdf_filename_format, 
                   retries, delay, dry_run=False, test_email=None):
    """Processes the admit card generation and email sending workflow.

    Args:
        data_file (str): Path to the Excel data file.
        template_file (str): Path to the .docx template file.
        pdf_folder (str): Path to the folder to save generated PDFs.
        logs_folder (str): Path to the folder to save the log file.
        log_callback (function): A function to call for logging messages.
        email_subject (str): The subject of the email.
        email_body (str): The HTML body of the email.
        pdf_filename_format (str): The format for the PDF filenames.
        retries (int): The number of times to retry sending an email.
        delay (int): The delay in seconds between emails.
        dry_run (bool, optional): If True, generates PDFs but doesn't send emails. Defaults to False.
        test_email (str, optional): If provided, sends a single test email. Defaults to None.
    """
    log_file_path = os.path.join(logs_folder, "documint_log.xlsx")
    
    if not os.access(data_file, os.R_OK):
        log_callback(f"❌ The Excel file '{data_file}' is not accessible. Please close it and try again.")
        return

    try:
        df = pd.read_excel(data_file)
        df.columns = df.columns.str.strip()
    except Exception as e:
        log_callback(f"❌ Error reading Excel file: {str(e)}")
        return

    outlook = win32com.client.Dispatch("Outlook.Application")
    log_records = []
    
    if test_email:
        # Create a dummy dataframe with all columns from the original df
        dummy_data = {col: f"Test {col}" for col in df.columns}
        dummy_data["Email"] = test_email
        df = pd.DataFrame([dummy_data])

    for index, row in df.iterrows():
        try:
            email = str(row["Email"]).strip()
            if not is_valid_email(email):
                log_records.append({
                    "Name": str(row.get("Name", "N/A")).strip(),
                    "Email": email,
                    "Status": f"Failed: Invalid Email Format",
                    "Timestamp": datetime.now()
                })
                log_callback(f"FAILED: Invalid Email Format for {email}")
                continue

            replacements = {f"<{col}>": str(val).strip() for col, val in row.items()}
            doc = Document(template_file)
            replace_placeholders_in_doc(doc, replacements)

            # Create filename from format string
            filename = pdf_filename_format.format(**replacements)
            word_filename = os.path.join(pdf_folder, f"{filename}.docx")
            pdf_filename = os.path.join(pdf_folder, f"{filename}.pdf")
            doc.save(word_filename)

            try:
                word_app = win32com.client.Dispatch("Word.Application")
                word_app.Visible = False
                word_doc = word_app.Documents.Open(os.path.abspath(word_filename))
                word_doc.SaveAs(os.path.abspath(pdf_filename), FileFormat=17)
                word_doc.Close(False)
                word_app.Quit()
            except Exception as conv_error:
                log_records.append({
                    "Name": str(row.get("Name", "N/A")).strip(),
                    "Email": email,
                    "Status": f"Failed: Word to PDF conversion error: {conv_error}",
                    "Timestamp": datetime.now()
                })
                log_callback(f"FAILED: Word to PDF conversion error for {email}")
                continue
            
            if not dry_run:
                sent = False
                for attempt in range(retries):
                    try:
                        mail = outlook.CreateItem(0)
                        mail.To = email
                        mail.Subject = email_subject
                        mail.HTMLBody = email_body.format(**replacements)
                        mail.Attachments.Add(os.path.abspath(pdf_filename))
                        mail.Send()
                        sent = True
                        break
                    except Exception as send_error:
                        if attempt == retries - 1:
                            log_records.append({
                                "Name": str(row.get("Name", "N/A")).strip(),
                                "Email": email,
                                "Status": f"Failed to send email: {str(send_error)}",
                                "Timestamp": datetime.now()
                            })
                            log_callback(f"FAILED: Could not send email to {email}")
                        else:
                            time.sleep(delay)

                if sent:
                    log_records.append({
                        "Name": str(row.get("Name", "N/A")).strip(),
                        "Email": email,
                        "Status": "Success",
                        "Timestamp": datetime.now()
                    })
                    log_callback(f"SUCCESS: Sent document to {email}")
            else:
                log_callback(f"DRY RUN: Generated PDF for {email}")
                log_records.append({
                    "Name": str(row.get("Name", "N/A")).strip(),
                    "Email": email,
                    "Status": "Dry Run - PDF Generated",
                    "Timestamp": datetime.now()
                })

            time.sleep(delay)

        except Exception as e:
            log_records.append({
                "Name": str(row.get("Name", "N/A")).strip(),
                "Email": str(row.get("Email", "N/A")).strip(),
                "Status": f"Failed to process row: {str(e)}",
                "Timestamp": datetime.now()
            })
            log_callback(f"FAILED: Could not process row {index + 2}: {str(e)}")

        finally:
            if os.path.exists(word_filename):
                os.remove(word_filename)

    try:
        if os.path.exists(log_file_path):
            old_log_df = pd.read_excel(log_file_path)
            new_log_df = pd.concat([old_log_df, pd.DataFrame(log_records)], ignore_index=True)
        else:
            new_log_df = pd.DataFrame(log_records)
        new_log_df.to_excel(log_file_path, index=False)
    except Exception as log_error:
        log_callback(f"❌ Could not save log file: {log_error}")